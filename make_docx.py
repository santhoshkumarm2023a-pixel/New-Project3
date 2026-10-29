import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def create_report():
    doc = Document()

    # Set standard margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Helper function to style headings
    def add_custom_heading(text, level=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Calibri"
        if level == 1:
            run.font.size = Pt(15)
            run.font.color.rgb = RGBColor(26, 54, 93)  # Navy
        elif level == 2:
            run.font.size = Pt(12.5)
            run.font.color.rgb = RGBColor(43, 108, 176)  # Blue
        return p

    # --- Header / Institution Block ---
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_after = Pt(2)
    r_inst = p_inst.add_run("LAB PROJECT REPORT SUBMISSION")
    r_inst.bold = True
    r_inst.font.size = Pt(11)
    r_inst.font.color.rgb = RGBColor(113, 128, 150)

    # --- Document Title ---
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(12)
    r_title = p_title.add_run("IoT-Based Smart RFID Authentication and Access Control Management System")
    r_title.bold = True
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(19)
    r_title.font.color.rgb = RGBColor(26, 54, 93)

    # --- Student Details Table ---
    tbl = doc.add_table(rows=4, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    metadata = [
        ("Student Name:", "Santhoshkumar M"),
        ("Register Number:", "23MID0080"),
        ("Course / Subject:", "CSI3007 - Advance Python Programming"),
        ("Submission Date:", "8th October 2026")
    ]

    for i, (label, val) in enumerate(metadata):
        row = tbl.rows[i]
        
        # Col 0
        c0 = row.cells[0]
        c0.width = Inches(2.2)
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(3)
        p0.paragraph_format.space_before = Pt(3)
        r0 = p0.add_run(label)
        r0.bold = True
        r0.font.size = Pt(10)
        r0.font.color.rgb = RGBColor(45, 55, 72)
        
        # Col 1
        c1 = row.cells[1]
        c1.width = Inches(4.3)
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(3)
        p1.paragraph_format.space_before = Pt(3)
        r1 = p1.add_run(val)
        r1.font.size = Pt(10)
        r1.font.color.rgb = RGBColor(26, 32, 44)

        # Style cell background
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F7FAFC"/>')
        c0._tc.get_or_add_tcPr().append(shading)
        shading2 = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F7FAFC"/>')
        c1._tc.get_or_add_tcPr().append(shading2)

    # Horizontal divider rule
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(10)
    p_div.paragraph_format.space_after = Pt(10)
    p_div_border = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:bottom w:val="single" w:sz="6" w:space="1" w:color="CBD5E0"/></w:pBdr>')
    p_div._p.get_or_add_pPr().append(p_div_border)

    # ==========================================
    # 1. Project Title
    # ==========================================
    add_custom_heading("1. Project Title", level=1)
    p1 = doc.add_paragraph()
    p1.paragraph_format.space_after = Pt(8)
    r = p1.add_run("IoT-Based Smart RFID Authentication and Access Control Management System")
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(43, 108, 176)

    # ==========================================
    # 2. What's your website do?
    # ==========================================
    add_custom_heading("2. What's your website do? (in 2 lines)", level=1)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(8)
    p2.paragraph_format.left_indent = Inches(0.2)
    
    # Left border highlight for quote
    quote_border = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:left w:val="single" w:sz="18" w:space="8" w:color="3182CE"/></w:pBdr>')
    p2._p.get_or_add_pPr().append(quote_border)
    
    r2 = p2.add_run(
        "A centralized security dashboard that authenticates RFID card scans from an Arduino reader in real time, granting or denying entry based on authorized credentials. "
        "It provides live access logs, card registration, and manual security controls for monitoring secure access points."
    )
    r2.italic = True
    r2.font.size = Pt(11)
    r2.font.color.rgb = RGBColor(45, 55, 72)

    # ==========================================
    # 3. Frontend [Screenshots]
    # ==========================================
    add_custom_heading("3. Frontend (Screenshots)", level=1)
    
    p_fe_intro = doc.add_paragraph()
    p_fe_intro.paragraph_format.space_after = Pt(6)
    r_fe_intro = p_fe_intro.add_run(
        "The frontend application is constructed using Python Streamlit, providing an interactive, modern, and responsive user interface for monitoring access control telemetry. Below are the functional screenshots captured from the live application:"
    )
    r_fe_intro.font.size = Pt(10.5)

    screenshots = [
        (
            "Figure 1: Main Overview & Live Access Logs Feed",
            "screenshots/screenshot1_live_access_feed.png",
            "This interface displays high-level security KPI metric cards (Total Scans Today, Access Granted, Access Denied, Registered Cards), the sidebar system connectivity indicators showing active hardware link to Arduino UNO (COM3) and the RC522 Reader, test card simulation controls, and the real-time access audit log table with granular timestamp, user, UID, role, status, and door location data."
        ),
        (
            "Figure 2: User & Card Registration Registry",
            "screenshots/screenshot2_user_card_registry.png",
            "This view features the credential enrollment interface. Administrators can onboard new personnel by entering their Full Name, scanning or inputting their RFID Card UID, and designating an access role (Student, Faculty, Security/Staff). Successfully registered credentials dynamically synchronize and appear immediately in the Authorized Users Database table."
        ),
        (
            "Figure 3: Access Analytics & Verification Breakdown",
            "screenshots/screenshot3_access_analytics.png",
            "This analytics dashboard visualizes system activity metrics via dynamic bar charts. The left chart depicts the Access Verification Breakdown (Granted vs. Denied scan distribution), while the right chart presents User Demographics by organizational role, enabling security personnel to evaluate door throughput and spot anomalies."
        )
    ]

    for title, img_path, desc in screenshots:
        add_custom_heading(title, level=2)
        
        # Add Image
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(4)
            p_img.paragraph_format.space_after = Pt(4)
            p_img.paragraph_format.keep_with_next = True
            run_img = p_img.add_run()
            run_img.add_picture(img_path, width=Inches(6.0))
        
        # Add Description
        p_desc = doc.add_paragraph()
        p_desc.paragraph_format.space_before = Pt(2)
        p_desc.paragraph_format.space_after = Pt(10)
        r_desc_label = p_desc.add_run("Description: ")
        r_desc_label.bold = True
        r_desc_label.font.size = Pt(10)
        r_desc = p_desc.add_run(desc)
        r_desc.font.size = Pt(10)
        r_desc.font.color.rgb = RGBColor(74, 85, 104)

    # ==========================================
    # 4. Backend Architecture
    # ==========================================
    add_custom_heading("4. Backend Architecture", level=1)
    
    p_arch_intro = doc.add_paragraph()
    p_arch_intro.paragraph_format.space_after = Pt(6)
    r_arch_intro = p_arch_intro.add_run(
        "The system follows a robust four-tier embedded IoT architecture designed for deterministic authentication latency and seamless communication between physical edge hardware and the Python web application:"
    )
    r_arch_intro.font.size = Pt(10.5)

    tiers = [
        ("1. Edge Hardware Layer: ", "Comprises the Arduino UNO microcontroller interfaced with the RC522 RFID reader module over the SPI (Serial Peripheral Interface) bus operating at 13.56 MHz, alongside output actuators (Status Green/Red LEDs, Active Buzzer, and optional Door Relay)."),
        ("2. Communication Bridge Layer: ", "Implements full-duplex asynchronous serial communication (UART over USB at 9600 baud rate) to stream captured tag UID packets from the edge microcontroller to the host workstation."),
        ("3. Application & Verification Layer: ", "The Python core backend (built with Python 3.14/3.x) decodes incoming UID serial packets, queries the authorized credentials database, evaluates role permissions, and returns immediate acknowledgment signals ('G' for Granted, 'D' for Denied) back to the microcontroller."),
        ("4. Presentation Layer: ", "Built with Python Streamlit, rendering real-time access telemetry, visual KPI metrics, enrollment management forms, and analytical audit charts.")
    ]

    for label, text in tiers:
        p_tier = doc.add_paragraph(style='List Bullet')
        p_tier.paragraph_format.space_before = Pt(2)
        p_tier.paragraph_format.space_after = Pt(3)
        r_lbl = p_tier.add_run(label)
        r_lbl.bold = True
        r_lbl.font.size = Pt(10)
        r_lbl.font.color.rgb = RGBColor(43, 108, 176)
        r_txt = p_tier.add_run(text)
        r_txt.font.size = Pt(10)
        r_txt.font.color.rgb = RGBColor(45, 55, 72)

    # Architecture ASCII Box
    p_box = doc.add_paragraph()
    p_box.paragraph_format.space_before = Pt(8)
    p_box.paragraph_format.space_after = Pt(8)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="EDF2F7"/>')
    p_box._p.get_or_add_pPr().append(shd)
    
    arch_ascii = (
        "[RFID Tag / Card]  -- (13.56 MHz RF) -->  [RC522 RFID Reader]\n"
        "                                                   │ (SPI Protocol: MOSI, MISO, SCK, SS)\n"
        "                                                   ▼\n"
        "                                           [Arduino UNO Edge Controller]\n"
        "                                                   │ (USB Serial / UART @ 9600 Baud)\n"
        "                                                   ▼\n"
        "                                        [Python Application Engine]\n"
        "                                                   │\n"
        "                          ┌────────────────────────┴────────────────────────┐\n"
        "                          ▼                                                 ▼\n"
        "             [Authorized Database / Registry]               [Streamlit Presentation UI]\n"
        "             (Session State & Verified Cards)                (Live Telemetry & Analytics)"
    )
    r_box = p_box.add_run(arch_ascii)
    r_box.font.name = "Consolas"
    r_box.font.size = Pt(8.5)
    r_box.font.color.rgb = RGBColor(45, 55, 72)

    # ==========================================
    # 5. Hosting Status
    # ==========================================
    add_custom_heading("5. Are you hosting? If yes, where are you hosting?", level=1)
    
    p5 = doc.add_paragraph()
    p5.paragraph_format.space_before = Pt(4)
    p5.paragraph_format.space_after = Pt(8)
    p5.paragraph_format.left_indent = Inches(0.2)
    
    border5 = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:left w:val="single" w:sz="18" w:space="8" w:color="3182CE"/></w:pBdr>')
    p5._p.get_or_add_pPr().append(border5)
    
    r5_bold = p5.add_run("Hosting Status: No, running on Localhost (http://localhost:8501).\n\n")
    r5_bold.bold = True
    r5_bold.font.size = Pt(11)
    r5_bold.font.color.rgb = RGBColor(26, 54, 93)

    r5_detail = p5.add_run(
        "Technical Justification: Since the authentication system directly interfaces with physical edge hardware—specifically the Arduino UNO microcontroller and RC522 RFID scanner via local hardware serial COM ports (UART over USB at COM3)—hosting locally ensures ultra-low-latency deterministic response times (<50ms) for physical door unlock solenoids, eliminates cloud roundtrip dependencies, and guarantees high-security offline operational capability during network outages."
    )
    r5_detail.font.size = Pt(10)
    r5_detail.font.color.rgb = RGBColor(45, 55, 72)

    output_path = "Project_Report_23MID0080.docx"
    doc.save(output_path)
    print(f"Successfully generated {output_path}")

if __name__ == "__main__":
    create_report()
